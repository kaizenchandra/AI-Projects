"""Bounded format validation. Text inputs retain original lexical values."""

import csv
import io
import zipfile
import os
from pathlib import Path
import openpyxl
import polars as pl

MAX_ROWS = int(os.getenv("EDA_MAX_ROWS", "100000"))
MAX_COLS = int(os.getenv("EDA_MAX_COLUMNS", "100"))
MAX_BYTES = int(os.getenv("EDA_MAX_BYTES", "20971520"))
if not (1 <= MAX_ROWS <= 1000000 and 1 <= MAX_COLS <= 500 and 1 <= MAX_BYTES <= 104857600):
    raise RuntimeError("Invalid configured dataset limits.")


def headers(values):
    names = [str(v).strip() if v is not None else "" for v in values]
    if any(not n for n in names):
        raise ValueError("Headers must be non-empty; use header=false for headerless files.")
    if len(set(names)) != len(names):
        raise ValueError("Duplicate column names are not allowed.")
    if len(names) > MAX_COLS:
        raise ValueError(f"Maximum {MAX_COLS} columns.")
    return names


def ingest(data: bytes, filename: str, delimiter=",", encoding="utf-8", header=True, sheet=None):
    if not data or len(data) > MAX_BYTES:
        raise ValueError("File is empty or exceeds the configured upload limit.")
    suffix = Path(filename).suffix.lower()
    sheets = []
    if suffix == ".parquet":
        if data[:4] != b"PAR1" or data[-4:] != b"PAR1":
            raise ValueError("Invalid Parquet signature.")
        # Lazy scan checks metadata before materializing potentially compressed data.
        scan = pl.scan_parquet(io.BytesIO(data))
        if len(scan.collect_schema()) > MAX_COLS or scan.select(pl.len()).collect().item() > MAX_ROWS:
            raise ValueError("Parquet exceeds configured dimensions.")
        df = scan.collect()
        if any(t.is_nested() or t == pl.Binary for t in df.dtypes):
            raise ValueError("Nested and binary Parquet columns are unsupported.")
        headers(df.columns)
    elif suffix == ".xlsx":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError("Invalid XLSX container.")
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024 or len(z.infolist()) > 2000:
                raise ValueError("Workbook decompression limit exceeded.")
            if any("vbaProject" in i.filename for i in z.infolist()):
                raise ValueError("Macro-enabled workbooks are unsupported.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = wb.sheetnames
            if sheet and sheet not in sheets:
                raise ValueError("Unknown worksheet.")
            ws = wb[sheet or sheets[0]]
            if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > MAX_COLS:
                raise ValueError("Workbook exceeds configured dimensions.")
            rows = []
            for row in ws.iter_rows(values_only=True):
                if len(rows) > MAX_ROWS:
                    raise ValueError("Row limit exceeded.")
                rows.append([None if v is None else str(v) for v in row])
            df = from_rows(rows, header)
        finally:
            wb.close()
    elif suffix == ".csv":
        if encoding not in ("utf-8", "utf-8-sig", "latin-1") or len(delimiter) != 1:
            raise ValueError("Unsupported encoding or delimiter.")
        text = data.decode(encoding, errors="strict")
        if "\x00" in text or data.startswith((b"PK", b"PAR1")):
            raise ValueError("File content is not CSV text.")
        rows = []
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True):
            if len(rows) > MAX_ROWS or len(row) > MAX_COLS:
                raise ValueError("Dataset exceeds configured dimensions.")
            rows.append([v if v != "" else None for v in row])
        df = from_rows(rows, header)
    else:
        raise ValueError("Supported formats: CSV, XLSX, Parquet.")
    if df.height == 0 or df.width == 0:
        raise ValueError("Dataset must contain at least one record and column.")
    if df.height > MAX_ROWS or df.width > MAX_COLS or df.estimated_size() > 150 * 1024 * 1024:
        raise ValueError("Dataset exceeds configured limits.")
    return df, sheets


def from_rows(rows, header):
    if not rows:
        raise ValueError("No records found.")
    names = headers(rows[0]) if header else [f"column_{i + 1}" for i in range(len(rows[0]))]
    values = rows[1:] if header else rows
    if any(len(r) != len(names) for r in values):
        raise ValueError("Malformed row: field count differs from header.")
    return pl.DataFrame(values, schema=names, orient="row", infer_schema_length=None)
